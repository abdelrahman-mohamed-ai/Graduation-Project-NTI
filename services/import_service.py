import io
import csv
import json
import re
import uuid
import zipfile
from datetime import datetime, timezone, timedelta
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from werkzeug.utils import secure_filename
from .model_service import ValidationError
from repositories.student_repository import now


def identity(data):
    uid = str(data.get('Student ID') or '').strip()
    name = str(data.get('Student Name') or '').strip()
    errors = {}
    if uid and not re.fullmatch(r'[A-Za-z0-9_-]{1,64}', uid): errors['Student ID'] = 'Use 1–64 letters, digits, hyphens or underscores.'
    if len(name) > 160: errors['Student Name'] = 'Use at most 160 characters.'
    if errors: raise ValidationError(errors)
    return uid, name


class ImportService:
    def __init__(self, model, repo): self.model, self.repo = model, repo

    def preview(self, upload, session_id):
        filename = secure_filename(upload.filename or '')
        extension = Path(filename).suffix.lower()
        if extension not in ('.csv', '.xlsx'): raise ValueError('Choose a CSV or XLSX file.')
        content = upload.read()
        if not content: raise ValueError('The file is empty.')
        try:
            if extension == '.xlsx':
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024: raise ValueError('Workbook expands beyond the processing limit.')
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                try: original_headers = list(next(workbook.active.iter_rows(values_only=True)))
                finally: workbook.close()
                frame = pd.read_excel(io.BytesIO(content), engine='openpyxl', dtype=object, nrows=5001)
            else:
                decoded = content.decode('utf-8-sig')
                dialect = csv.Sniffer().sniff(decoded[:65536], delimiters=',;\t') if any(c in decoded.splitlines()[0] for c in ',;\t') else csv.excel
                original_headers = next(csv.reader(io.StringIO(decoded), dialect))
                frame = pd.read_csv(io.BytesIO(content), sep=None, engine='python', dtype=object, nrows=5001, encoding='utf-8-sig', keep_default_na=False)
        except Exception as exc: raise ValueError('The file could not be read. Use a valid UTF-8 CSV or an unencrypted XLSX workbook.') from exc
        if not len(frame): raise ValueError('The file contains no student rows.')
        normalized_headers = [str(v).strip() for v in original_headers]
        if len(set(normalized_headers)) != len(normalized_headers): raise ValueError('Column names must be unique after trimming whitespace.')
        if len(frame) > 5000: raise ValueError('Import up to 5,000 students per file.')
        frame.columns = frame.columns.astype(str).str.strip()
        if frame.columns.duplicated().any(): raise ValueError('Column names must be unique after trimming whitespace.')
        frame = frame.where(pd.notna(frame), None)
        required = self.model.features
        missing = [f for f in required if f not in frame.columns]
        extras = [f for f in frame.columns if f not in required and f not in ('Student ID', 'Student Name')]
        rows, problems = [], []
        existing = {s['student_uid'] for s in self.repo.query(all_rows=True)['students']}
        # Mark every occurrence of a repeated supplied ID; do not pick an arbitrary row.
        counts = frame['Student ID'].fillna('').astype(str).str.strip().value_counts() if 'Student ID' in frame else {}
        for i, row in enumerate(frame.to_dict('records'), 2):
            try:
                if missing: raise ValidationError({'columns': 'Required columns are missing.'})
                uid, name = identity(row)
                if uid and (counts.get(uid, 0) > 1 or uid in existing): raise ValidationError({'Student ID': 'Duplicate ID in this file or the database. No existing student will be overwritten.'})
                features = self.model.validate(row)
                rows.append({'uid': uid, 'name': name, 'features': features, 'row': i})
            except ValidationError as exc: problems.append({'row': i, 'errors': exc.errors})
        token = uuid.uuid4().hex
        payload = {'rows': rows, 'problems': problems, 'row_count': len(frame)}
        with self.repo.connection() as db:
            db.execute('DELETE FROM import_batches WHERE created_at < ?', ((datetime.now(timezone.utc) - timedelta(days=1)).isoformat(),))
            db.execute('INSERT INTO import_batches(token,session_id,filename,payload_json,created_at) VALUES(?,?,?,?,?)', (token, session_id, filename, json.dumps(payload), now()))
        return {'token': token, 'filename': filename, 'size': len(content), 'row_count': len(frame), 'column_count': len(frame.columns),
            'columns': list(frame.columns), 'preview': json.loads(frame.head(10).to_json(orient='records')),
            'required_found': [f for f in required if f in frame], 'missing': missing, 'extra': extras,
            'problems': problems, 'valid_rows': len(rows)}

    def commit(self, token, session_id):
        # Claim the batch before model work; a repeated click cannot duplicate students.
        with self.repo.connection() as db:
            db.execute('BEGIN IMMEDIATE')
            batch = db.execute('SELECT * FROM import_batches WHERE token=? AND session_id=?', (token, session_id)).fetchone()
            if not batch: raise ValueError('Upload preview expired. Select the file again.')
            if batch['status'] == 'complete': return json.loads(batch['result_json'])
            if batch['status'] != 'preview': raise ValueError('This import is already processing.')
            data = json.loads(batch['payload_json'])
            if not data['rows']: raise ValueError('There are no valid rows to import.')
            db.execute("UPDATE import_batches SET status='processing' WHERE token=?", (token,))
        try:
            results = self.model.predict_many([r['features'] for r in data['rows']])
            summary = {'imported': 0, 'analyzed': 0, 'High': 0, 'Medium': 0, 'Low': 0, 'failed_rows': data['problems'], 'student_uids': [], 'token': token}
            with self.repo.connection() as db:
                db.execute('BEGIN IMMEDIATE')
                # Explicit IDs are inserted first so generated IDs cannot consume them.
                for row, result in sorted(zip(data['rows'], results), key=lambda pair: not bool(pair[0]['uid'])):
                    if row['uid'] and db.execute('SELECT 1 FROM students WHERE student_uid=?', (row['uid'],)).fetchone():
                        summary['failed_rows'].append({'row': row['row'], 'errors': {'Student ID': 'ID was imported since preview. Row skipped.'}}); continue
                    uid = self.repo.insert_result(db, row['uid'], row['name'], 'Upload', result)
                    summary['student_uids'].append(uid)
                    summary['imported'] += 1; summary['analyzed'] += 1; summary[result['risk_level']] += 1
                db.execute('INSERT INTO data_imports(filename,row_count,success_count,failed_count,created_at) VALUES(?,?,?,?,?)',
                    (batch['filename'], data['row_count'], summary['imported'], len(summary['failed_rows']), now()))
                db.execute("UPDATE import_batches SET status='complete',result_json=? WHERE token=?", (json.dumps(summary), token))
            return summary
        except Exception:
            with self.repo.connection() as db: db.execute("UPDATE import_batches SET status='preview' WHERE token=?", (token,))
            raise
